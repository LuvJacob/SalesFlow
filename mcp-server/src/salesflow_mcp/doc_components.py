"""Reusable python-docx renderer components for SalesFlow Document Studio.

Every document template is composed from the functions in this module so that
no template duplicates styling code: title block, headings, metric cards,
comparison matrix, risk table, action table, source references, evidence
legend and disclaimer.

Design: restrained navy / white / blue palette, no logos, readability first.
Colour is never the only signal - every indicator also carries a text label.

This module formats supplied content only. It never invents facts.
"""

from __future__ import annotations

import re
from typing import Any, Iterable, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1A, 0x3A, 0x6C)
BLUE = RGBColor(0x1F, 0x5B, 0xAC)
MUTED = RGBColor(0x57, 0x60, 0x6A)
INK = RGBColor(0x1F, 0x23, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADER_BG = "1A3A6C"
ROW_ALT_BG = "F3F6FB"
CARD_BG = "EBF0F8"
BORDER = "C9D3E3"

FONT = "Calibri"
CONTENT_WIDTH_IN = 6.9  # Letter, 0.8" side margins

# Semantic indicator palette: (fill hex, text colour)
_GREEN = ("E3F1E6", RGBColor(0x1E, 0x6B, 0x34))
_AMBER = ("FFF3D6", RGBColor(0x8A, 0x5A, 0x00))
_RED = ("FBE3E3", RGBColor(0x9B, 0x1C, 0x1C))
_GREY = ("ECEEF1", RGBColor(0x57, 0x60, 0x6A))
_BLUE = ("E3ECF8", RGBColor(0x1F, 0x5B, 0xAC))

FIT_STYLES = {
    "supported": _GREEN,
    "partially supported": _AMBER,
    "not supported": _RED,
    "unknown / insufficient evidence": _GREY,
}
RISK_STYLES = {"high": _RED, "medium": _AMBER, "low": _GREEN}

# Information-integrity labels used across every document.
LABEL_VERIFIED = "Verified"
LABEL_SYNTHETIC = "Synthetic CRM"
LABEL_AI = "AI-generated"
LABEL_UNKNOWN = "Unknown"
LABEL_UNVERIFIED = "Unverified spec"

LABEL_STYLES = {
    LABEL_VERIFIED: _GREEN,
    LABEL_SYNTHETIC: _BLUE,
    LABEL_AI: _AMBER,
    LABEL_UNKNOWN: _GREY,
    LABEL_UNVERIFIED: _RED,
}

LEGEND_ROWS = [
    (LABEL_VERIFIED, "Confirmed by a cited source that the seller can check."),
    (LABEL_SYNTHETIC, "Fictional demonstration CRM data; not a real customer record."),
    (LABEL_AI, "Produced by the assistant (a recommendation or an extraction) from the data shown; needs seller judgement."),
    (LABEL_UNKNOWN, "Not present in the data. Must be confirmed before relying on it."),
    (LABEL_UNVERIFIED, "Product capability from an unverified or placeholder catalog entry."),
]

# ---------------------------------------------------------------------------
# Text hygiene
# ---------------------------------------------------------------------------
# XML 1.0 forbids most control characters; python-docx raises on them.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def clean(value: Any) -> str:
    """Return text safe for OOXML (no control characters). None -> ''."""
    if value is None:
        return ""
    return _CONTROL.sub("", str(value))


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------


# Schema element order matters: Word tolerates violations, strict validators
# do not. Each helper inserts before the first successor that is present.
_TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_TCPR_AFTER_MAR = ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
_PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                   "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                   "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                   "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                   "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_PPR_AFTER_TABS = _PPR_AFTER_PBDR[2:]


def _put(parent, child, successors) -> None:
    """Insert ``child`` in schema order, replacing any existing element of the same tag."""
    for old in parent.findall(child.tag):
        parent.remove(old)
    parent.insert_element_before(child, *successors)


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    _put(tcPr, shd, _TCPR_AFTER_SHD)


def _cell_margins(cell, top=50, bottom=50, left=90, right=90) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _put(tcPr, mar, _TCPR_AFTER_MAR)


def _table_borders(table, color: str = BORDER) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    _put(tblPr, borders, _TBLPR_AFTER_BORDERS)


def _row_flags(row, header: bool = False) -> None:
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)
    if header:
        th = OxmlElement("w:tblHeader")
        trPr.append(th)


def _fixed_layout(table, widths_in: Sequence[float]) -> None:
    table.autofit = False  # python-docx writes <w:tblLayout w:type="fixed"/>
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is not None:
        # Percentage width (5000 = 100%) is honoured by Word and by lighter
        # renderers; column widths below still fix the proportions.
        pct = int(round(sum(widths_in) / CONTENT_WIDTH_IN * 5000))
        tblW.set(qn("w:type"), "pct")
        tblW.set(qn("w:w"), str(min(pct, 5000)))
    for row in table.rows:
        for idx, w in enumerate(widths_in):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(w)
    grid = table._tbl.tblGrid
    for idx, col in enumerate(grid.findall(qn("w:gridCol"))):
        if idx < len(widths_in):
            col.set(qn("w:w"), str(int(widths_in[idx] * 1440)))


def _fld_simple(paragraph, instr: str, size_pt: float, color: RGBColor) -> None:
    """Append a PAGE / NUMPAGES field as begin/instr/separate/result/end runs.

    Each run carries explicit formatting so the updated field result keeps the
    footer's size and colour (a bare fldSimple falls back to the style size).
    """

    def run_with(child):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color")
        col.set(qn("w:val"), str(color))
        rpr.append(col)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size_pt * 2)))
        rpr.append(sz)
        r.append(rpr)
        r.append(child)
        paragraph._p.append(r)

    def fld_char(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run_with(fld_char("begin"))
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run_with(instr_el)
    run_with(fld_char("separate"))
    t = OxmlElement("w:t")
    t.text = "1"
    run_with(t)
    run_with(fld_char("end"))


def _bottom_border(paragraph, color: str = HEADER_BG, size: int = 8) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    _put(pPr, pbdr, _PPR_AFTER_PBDR)


def _top_border(paragraph, color: str = BORDER, size: int = 4) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), color)
    pbdr.append(top)
    _put(pPr, pbdr, _PPR_AFTER_PBDR)


def _run(paragraph, text: str, *, bold=False, italic=False, color: Optional[RGBColor] = None,
         size: Optional[float] = None):
    r = paragraph.add_run(clean(text))
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    return r


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------


def _style_font(style, size: float, *, bold=False, color: RGBColor = INK, italic=False) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    # Drop theme font overrides so the explicit font wins.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def new_document(*, title: str, header_note: str, footer_note: str) -> Document:
    """Create a Letter-size document with styles, header and footer.

    The footer carries ``footer_note`` and automatic "Page X of Y" fields.
    """
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.8)
    sec.header_distance = Inches(0.4)
    sec.footer_distance = Inches(0.4)

    styles = doc.styles
    _style_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    _style_font(styles["Title"], 24, bold=True, color=NAVY)
    styles["Title"].paragraph_format.space_after = Pt(2)
    # Remove the default Title bottom border from the template.
    t_ppr = styles["Title"].element.get_or_add_pPr()
    for b in t_ppr.findall(qn("w:pBdr")):
        t_ppr.remove(b)
    _style_font(styles["Heading 1"], 15, bold=True, color=NAVY)
    styles["Heading 1"].paragraph_format.space_before = Pt(16)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    _style_font(styles["Heading 2"], 12, bold=True, color=BLUE)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    _style_font(styles["Heading 3"], 10.5, bold=True, color=INK)
    styles["Heading 3"].paragraph_format.space_before = Pt(6)
    styles["Heading 3"].paragraph_format.space_after = Pt(2)
    styles["Heading 3"].paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        _style_font(styles[name], 10.5)
        styles[name].paragraph_format.space_after = Pt(2)

    # Header/Footer styles ship with centre and right tab stops; remove them so
    # the only tab stop is the right-aligned one set on each paragraph.
    for name in ("Header", "Footer"):
        ppr = styles[name].element.get_or_add_pPr()
        for t in ppr.findall(qn("w:tabs")):
            ppr.remove(t)

    # Header: report title (left) and draft marker (right)
    hp = sec.header.paragraphs[0]
    hp.text = ""
    _set_right_tab(hp)
    _run(hp, title, size=8.5, color=MUTED)
    _run(hp, "\t" + header_note, size=8.5, color=MUTED, italic=True)
    _bottom_border(hp, color=BORDER, size=4)

    # Footer: note (left) and Page X of Y (right)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    _set_right_tab(fp)
    _top_border(fp)
    _run(fp, footer_note, size=8, color=MUTED)
    _run(fp, "\tPage ", size=8.5, color=MUTED)
    _fld_simple(fp, "PAGE", 8.5, MUTED)
    _run(fp, " of ", size=8.5, color=MUTED)
    _fld_simple(fp, "NUMPAGES", 8.5, MUTED)

    doc.core_properties.title = clean(title)[:250]
    doc.core_properties.author = "SalesFlow Document Studio"
    return doc


def _set_right_tab(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(CONTENT_WIDTH_IN * 1440)))
    tabs.append(tab)
    _put(pPr, tabs, _PPR_AFTER_TABS)


# ---------------------------------------------------------------------------
# Text blocks
# ---------------------------------------------------------------------------


def add_title_block(doc, *, title: str, subtitle: str = "", metadata: Iterable[tuple[str, str]] = ()) -> None:
    """Report title, optional subtitle, and a compact metadata line table."""
    p = doc.add_paragraph(style="Title")
    _run(p, title)
    if subtitle:
        sp = doc.add_paragraph()
        _run(sp, subtitle, size=12.5, color=BLUE)
        sp.paragraph_format.space_after = Pt(4)
    rule = doc.add_paragraph()
    _bottom_border(rule, size=12)
    rule.paragraph_format.space_after = Pt(4)
    meta = [(clean(k), clean(v)) for k, v in metadata if clean(v)]
    if meta:
        tbl = doc.add_table(rows=0, cols=2)
        _table_borders_none(tbl)
        for k, v in meta:
            cells = tbl.add_row().cells
            _cell_write(cells[0], k, bold=True, color=MUTED, size=9.5)
            _cell_write(cells[1], v, size=9.5)
        _fixed_layout(tbl, [1.5, CONTENT_WIDTH_IN - 1.5])
        for row in tbl.rows:
            _row_flags(row)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _table_borders_none(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    _put(tblPr, borders, _TBLPR_AFTER_BORDERS)


def add_h1(doc, text: str) -> None:
    doc.add_paragraph(clean(text), style="Heading 1")


def add_h2(doc, text: str) -> None:
    doc.add_paragraph(clean(text), style="Heading 2")


def add_h3(doc, text: str) -> None:
    doc.add_paragraph(clean(text), style="Heading 3")


def add_text(doc, text: str, *, muted=False, italic=False, size: Optional[float] = None) -> None:
    p = doc.add_paragraph()
    _run(p, text, italic=italic, color=MUTED if muted else None, size=size)


def add_labelled_text(doc, label: str, text: str) -> None:
    """Paragraph prefixed with an integrity label tag (text, not colour only)."""
    p = doc.add_paragraph()
    _label_run(p, label)
    _run(p, "  " + text)
    _keep_with_next(p)  # a label line introduces what follows; never strand it at a page end


def add_bullets(doc, items: Iterable[Any], *, numbered=False, empty_text: Optional[str] = None) -> None:
    values = [clean(i) for i in items if clean(i).strip()]
    if not values:
        if empty_text:
            add_text(doc, empty_text, muted=True, italic=True)
        return
    style = "List Number" if numbered else "List Bullet"
    for v in values:
        doc.add_paragraph(v, style=style)


def _label_run(paragraph, label: str):
    fill_text = LABEL_STYLES.get(label, _GREY)[1]
    return _run(paragraph, f"[{label.upper()}]", bold=True, color=fill_text, size=8.5)


def add_not_provided(doc, what: str) -> None:
    p = doc.add_paragraph()
    _label_run(p, LABEL_UNKNOWN)
    _run(p, f"  {clean(what)}: not provided. Confirm with the customer or CRM.", italic=True, color=MUTED)


# ---------------------------------------------------------------------------
# Table cells
# ---------------------------------------------------------------------------


def _cell_write(cell, content: Any, *, bold=False, italic=False, color: Optional[RGBColor] = None,
                size: float = 9.5, align: Optional[WD_ALIGN_PARAGRAPH] = None) -> None:
    """Write text (or a list of lines rendered as bullets) into a table cell."""
    cell.text = ""
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(0)
    first.paragraph_format.line_spacing = 1.05
    _cell_margins(cell)
    if isinstance(content, (list, tuple)):
        lines = [clean(c) for c in content if clean(c).strip()]
        if not lines:
            lines = ["-"]
        for i, line in enumerate(lines):
            para = first if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.05
            _run(para, ("\u2022 " if len(lines) > 1 else "") + line, bold=bold, italic=italic,
                 color=color, size=size)
            if align is not None:
                para.alignment = align
    else:
        _run(first, clean(content) or "-", bold=bold, italic=italic, color=color, size=size)
        if align is not None:
            first.alignment = align


def _pill(cell, text: str, style: tuple, *, size: float = 9) -> None:
    fill, fg = style
    _shade(cell, fill)
    _cell_write(cell, text, bold=True, color=fg, size=size, align=WD_ALIGN_PARAGRAPH.CENTER)


def _header_row(table, headers: Sequence[str]) -> None:
    cells = table.rows[0].cells
    for c, h in zip(cells, headers):
        _shade(c, HEADER_BG)
        _cell_write(c, h, bold=True, color=WHITE, size=9.5)
    _row_flags(table.rows[0], header=True)


def _data_table(doc, headers: Sequence[str], widths: Sequence[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(table)
    _header_row(table, headers)
    return table


KEEP_TOGETHER_MAX_ROWS = 14


def _finish_table(table, widths: Sequence[float], keep_together: bool = True) -> None:
    rows = table.rows
    if keep_together and len(rows) <= KEEP_TOGETHER_MAX_ROWS:
        # Small tables stay on one page: every row but the last keeps with the next.
        for row in list(rows)[:-1]:
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.keep_with_next = True
    for i, row in enumerate(table.rows[1:]):
        _row_flags(row)
        if i % 2 == 1:
            for c in row.cells:
                if c._tc.tcPr is None or c._tc.tcPr.find(qn("w:shd")) is None:
                    _shade(c, ROW_ALT_BG)
    _fixed_layout(table, widths)
    spacer = table._tbl.getnext()
    # Add a small paragraph after the table so following content has breathing room.
    p = OxmlElement("w:p")
    if spacer is not None:
        spacer.addprevious(p)
    else:
        table._tbl.addnext(p)


def _scale(widths: Sequence[float]) -> list[float]:
    total = sum(widths)
    return [w * CONTENT_WIDTH_IN / total for w in widths]


# ---------------------------------------------------------------------------
# Composite components
# ---------------------------------------------------------------------------


def add_metric_cards(doc, metrics: Sequence[tuple[str, Any]]) -> None:
    """A single row of 'cards': large value over a small label."""
    metrics = [(clean(l), clean(v)) for l, v in metrics]
    if not metrics:
        return
    n = len(metrics)
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(table, color="FFFFFF")
    for cell, (label, value) in zip(table.rows[0].cells, metrics):
        _shade(cell, CARD_BG)
        cell.text = ""
        _cell_margins(cell, top=90, bottom=90, left=110, right=110)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _run(p, value or "-", bold=True, size=17, color=NAVY)
        lp = cell.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_after = Pt(0)
        _run(lp, label, size=8.5, color=MUTED)
    _row_flags(table.rows[0])
    _fixed_layout(table, [CONTENT_WIDTH_IN / n] * n)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_key_value_table(doc, rows: Sequence[tuple[str, Any]]) -> None:
    rows = [(clean(k), v) for k, v in rows]
    if not rows:
        return
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(table)
    for k, v in rows:
        cells = table.add_row().cells
        _shade(cells[0], CARD_BG)
        _cell_write(cells[0], k, bold=True, color=NAVY)
        _cell_write(cells[1], v)
    for row in table.rows:
        _row_flags(row)
    _fixed_layout(table, [1.8, CONTENT_WIDTH_IN - 1.8])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[Any]], widths: Sequence[float]) -> None:
    """Generic data table with a repeating navy header row."""
    if not rows:
        return
    w = _scale(widths)
    table = _data_table(doc, headers, w)
    for r in rows:
        cells = table.add_row().cells
        for c, val in zip(cells, r):
            _cell_write(c, val)
    _finish_table(table, w)


def add_risk_table(doc, deals: Sequence[dict]) -> None:
    """Ranked risk table. Each deal: rank, name, stage, rating, score, evidence[], actions[]."""
    if not deals:
        return
    headers = ["#", "Deal / stage", "Risk", "Evidence", "Recommended action (AI)"]
    w = _scale([0.35, 1.5, 0.85, 2.3, 2.1])
    table = _data_table(doc, headers, w)
    for d in deals:
        cells = table.add_row().cells
        _cell_write(cells[0], d.get("rank", ""), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_write(cells[1], clean(d.get("name", "")) or "-", bold=True)
        for extra, muted in ((d.get("stage"), False), (d.get("value_text"), True)):
            if clean(extra):
                para = cells[1].add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                _run(para, clean(extra), size=9 if muted else 9.5, color=MUTED if muted else None)
        rating = clean(d.get("rating", "")) or "Unrated"
        style = RISK_STYLES.get(rating.split()[0].lower(), _GREY)
        score = d.get("score")
        text = f"{rating}\nScore {score}" if score not in (None, "") else rating
        _pill(cells[2], text, style)
        _cell_write(cells[3], list(d.get("evidence") or []) or ["No evidence recorded."])
        _cell_write(cells[4], list(d.get("actions") or []) or ["No action recorded."])
    _finish_table(table, w, keep_together=False)


def add_action_table(doc, actions: Sequence[dict]) -> None:
    """Action items with owner, deadline and basis. Missing values are shown as such."""
    if not actions:
        return
    headers = ["#", "Action", "Owner", "Deadline", "Basis"]
    w = _scale([0.35, 3.1, 1.15, 1.1, 1.1])
    table = _data_table(doc, headers, w)
    for i, a in enumerate(actions, 1):
        cells = table.add_row().cells
        _cell_write(cells[0], i, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_write(cells[1], a.get("action", ""))
        owner = clean(a.get("owner"))
        deadline = clean(a.get("deadline"))
        _cell_write(cells[2], owner or "Unassigned", italic=not owner, color=None if owner else MUTED)
        _cell_write(cells[3], deadline or "Not set", italic=not deadline, color=None if deadline else MUTED)
        basis = clean(a.get("basis")) or LABEL_AI
        _pill(cells[4], basis, LABEL_STYLES.get(basis, _GREY), size=8.5)
    _finish_table(table, w)


def add_comparison_matrix(doc, *, row_header: str, columns: Sequence[str],
                          rows: Sequence[tuple[str, Sequence[str]]]) -> None:
    """Requirement x product matrix; each cell is a fit status with colour + text."""
    if not rows or not columns:
        return
    headers = [row_header] + [clean(c) for c in columns]
    first_w = 2.6
    other = (CONTENT_WIDTH_IN - first_w) / len(columns)
    w = [first_w] + [other] * len(columns)
    table = _data_table(doc, headers, w)
    for label, statuses in rows:
        cells = table.add_row().cells
        _cell_write(cells[0], label)
        for c, st in zip(cells[1:], statuses):
            st = clean(st) or "Unknown / insufficient evidence"
            _pill(c, st, FIT_STYLES.get(st.lower(), _GREY), size=8.5)
    _finish_table(table, w)


def add_source_table(doc, sources: Sequence[dict]) -> None:
    """Source references: id, title, type, note. Preserved exactly as supplied."""
    if not sources:
        return
    headers = ["ID", "Source", "Type", "Note"]
    w = _scale([1.2, 2.4, 1.1, 2.4])
    table = _data_table(doc, headers, w)
    for s in sources:
        cells = table.add_row().cells
        _cell_write(cells[0], s.get("source_id", ""), size=9)
        title = clean(s.get("title") or s.get("name") or "")
        url = clean(s.get("url"))
        _cell_write(cells[1], title + (f"\n{url}" if url else "") or "-", size=9)
        _cell_write(cells[2], s.get("source_type", ""), size=9)
        _cell_write(cells[3], s.get("note", ""), size=9)
    _finish_table(table, w)


def add_evidence_legend(doc) -> None:
    """Explains the integrity labels used in this document."""
    table = _data_table(doc, ["Label", "Meaning"], [1.7, CONTENT_WIDTH_IN - 1.7])
    for label, meaning in LEGEND_ROWS:
        cells = table.add_row().cells
        _pill(cells[0], label, LABEL_STYLES[label], size=8.5)
        _cell_write(cells[1], meaning, size=9)
    _finish_table(table, [1.7, CONTENT_WIDTH_IN - 1.7])


def add_disclaimer(doc, text: str) -> None:
    """Shaded disclaimer box (single-cell table) - text only, no decorative icons."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_borders(table)
    cell = table.rows[0].cells[0]
    _shade(cell, "F7F8FA")
    _cell_write(cell, text, italic=True, color=MUTED, size=9)
    _row_flags(table.rows[0])
    _fixed_layout(table, [CONTENT_WIDTH_IN])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
